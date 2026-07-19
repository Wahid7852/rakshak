#!/usr/bin/env python3
# Builds RAKSHAK-deck.pptx from the same content/order as deck.html, so there's an
# editable native slide deck alongside the HTML one. Run: python generate_pptx.py
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pathlib import Path

PAPER = RGBColor(0xF4, 0xF1, 0xEA)
INK = RGBColor(0x24, 0x1F, 0x16)
MUTED = RGBColor(0x6B, 0x62, 0x52)
HAIRLINE = RGBColor(0xC7, 0xBC, 0xA2)
RUST = RGBColor(0xB5, 0x50, 0x2A)
QUANTUM = RGBColor(0x4B, 0x3F, 0x7A)
SURFACE = RGBColor(0xEA, 0xE4, 0xD4)

SERIF = "Georgia"
SANS = "Calibri"
MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def new_slide():
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    return slide


def textbox(slide, left, top, width, height, text, size=14, color=INK, font=SANS,
            bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=None, line_spacing=1.15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = font
        r.font.bold = bold
        r.font.italic = italic
    return box


def hline(slide, left, top, width, color=HAIRLINE):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(1))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()
    box.shadow.inherit = False
    return box


def chip(slide, left, top, width, height, text, size=11, border=HAIRLINE, text_color=INK,
          fill=SURFACE, dashed=False):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = border
    box.line.width = Pt(1)
    if dashed:
        ln = box.line._get_or_add_ln()
        prstDash = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(prstDash)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = MONO
    r.font.color.rgb = text_color
    return box


MARGIN = Inches(0.55)
CONTENT_W = SLIDE_W - 2 * MARGIN


def eyebrow(slide, label, num):
    textbox(slide, MARGIN, Inches(0.35), Inches(4), Inches(0.35), label.upper(),
            size=11, color=MUTED, font=MONO)
    textbox(slide, SLIDE_W - MARGIN - Inches(2), Inches(0.35), Inches(2), Inches(0.35), num,
            size=11, color=MUTED, font=MONO, align=PP_ALIGN.RIGHT)
    hline(slide, MARGIN, Inches(0.75), CONTENT_W)


def heading(slide, text, top=Inches(0.95), size=28):
    textbox(slide, MARGIN, top, CONTENT_W, Inches(0.7), text, size=size, color=INK,
            font=SERIF, bold=True)


def body_slide(label, num, title):
    slide = new_slide()
    eyebrow(slide, label, num)
    heading(slide, title)
    return slide


# 0. Title -----------------------------------------------------------------
slide = new_slide()
textbox(slide, MARGIN, Inches(1.6), Inches(9), Inches(0.4),
        "THEME 2 · QUANTUM MACHINE LEARNING FOR THREAT DETECTION",
        size=13, color=QUANTUM, font=MONO)
textbox(slide, MARGIN, Inches(2.1), Inches(10), Inches(1.3), "RAKSHAK",
        size=64, color=INK, font=SERIF, bold=True)
textbox(slide, MARGIN, Inches(3.3), Inches(9), Inches(0.9),
        "Quantum Malware Hunter: a hybrid classical and quantum cascade for\n"
        "real-time malware and behavior-log threat detection.",
        size=17, color=MUTED, font=SERIF, italic=True)
hline(slide, MARGIN, Inches(6.6), CONTENT_W)
textbox(slide, MARGIN, Inches(6.75), Inches(9), Inches(0.4),
        "Garima Shrivastava  ·  Abdul Wahid Khan  ·  Anna Mariya Martin",
        size=11, color=MUTED, font=MONO)
textbox(slide, SLIDE_W - MARGIN - Inches(4.5), Inches(6.75), Inches(4.5), Inches(0.4),
        "Local-first, HTTP + gRPC, Qt6 desktop console",
        size=11, color=MUTED, font=MONO, align=PP_ALIGN.RIGHT)

# 1. Problem statement -------------------------------------------------------
slide = body_slide("Problem", "01 / 12", "Problem statement")
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.05), Pt(3), Inches(1.9))
box.fill.solid(); box.fill.fore_color.rgb = RUST; box.line.fill.background(); box.shadow.inherit = False
textbox(slide, MARGIN + Inches(0.25), Inches(1.05), CONTENT_W - Inches(0.25), Inches(1.5),
        "Use quantum-enhanced ML to detect unknown malware patterns in real time. "
        "Build a hybrid AI system, classical and quantum, that predicts ransomware "
        "deployment patterns from user behavior logs.",
        size=15, color=INK, font=SERIF, italic=True)
textbox(slide, MARGIN + Inches(0.25), Inches(2.55), CONTENT_W - Inches(0.25), Inches(0.4),
        "Theme 2, “Quantum Malware Hunter & AI + Quantum = Early Ransomware Detection”",
        size=10, color=MUTED, font=MONO)
textbox(slide, MARGIN, Inches(3.2), CONTENT_W, Inches(1.5),
        "The gap: signature-based tools only catch what they've already seen. "
        "Classical-only ML plateaus on borderline, never-seen cases. Most “quantum "
        "security” demos are quantum-only toy models, with no classical backbone, no "
        "real-time constraint, no system an operator could actually run.",
        size=14, color=INK, font=SANS)
textbox(slide, MARGIN, Inches(4.7), CONTENT_W, Inches(1.6),
        "Ransomware has a shape before it detonates: mass file touches, auth anomalies, "
        "privilege-escalation attempts, all visible in ordinary behavior logs, minutes "
        "before the payload executes. The right answer scores that stream in real time "
        "and knows when to escalate an ambiguous case to a quantum judgment.",
        size=13, color=MUTED, font=SANS)

# 2. Solution overview (table) ----------------------------------------------
slide = body_slide("Solution", "02 / 12", "Solution overview")
rows = [
    ("Hybrid AI system: classical + quantum",
     "file_qsvc, a PennyLane quantum-embedded SVM, fused with classical detectors via "
     "confidence-weighted averaging. A real cascade stage, not a side demo."),
    ("Detect unknown malware patterns",
     "log_hst / log_ngram calibrate online with no labels; file_static / file_ml_or_rf "
     "score learned structural features, not signatures."),
    ("Predict ransomware from behavior logs",
     "The log cascade scores arbitrary log streams for anomalies in real time. Honestly: "
     "no ransomware-labeled model yet, positioned for it, not claiming it."),
    ("Real time",
     "Latency-budgeted cascade: 5ms p95 per log event, 50ms for an initial file verdict, "
     "250ms total with sandbox enrichment."),
    ("Dataset: simulated traffic or signatures",
     "Real data instead: LogHub HDFS_v1 and EMBER2018, a deliberate upgrade over the "
     "suggested baseline."),
]
tbl_top = Inches(1.05)
tbl = slide.shapes.add_table(len(rows), 2, MARGIN, tbl_top, CONTENT_W, Inches(5.6)).table
tbl.columns[0].width = Inches(4.2)
tbl.columns[1].width = CONTENT_W - Inches(4.2)
for i, (ask, ships) in enumerate(rows):
    c0, c1 = tbl.cell(i, 0), tbl.cell(i, 1)
    c0.text = ask
    c1.text = ships
    for cell, size, color, bold in ((c0, 12, MUTED, False), (c1, 12.5, INK, False)):
        cell.fill.solid(); cell.fill.fore_color.rgb = PAPER
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = SANS; r.font.bold = bold

# 3. System architecture -----------------------------------------------------
slide = body_slide("Architecture", "03 / 12", "System architecture")
cw, ch = Inches(9.5), Inches(0.5)
cx = MARGIN + (CONTENT_W - cw) / 2
y = Inches(1.15)
chip(slide, cx, y, cw, ch, "Qt6 desktop client: dashboard, scan, log analysis, quarantine, settings")
y += ch + Inches(0.12)
half = Inches(4.6)
chip(slide, cx, y, half, ch, "FastAPI (REST)")
chip(slide, cx + cw - half, y, half, ch, "gRPC Decider service")
y += ch + Inches(0.12)
chip(slide, cx, y, cw, ch, "Auth and rate limits: x-api-key header")
y += ch + Inches(0.12)
chip(slide, cx, y, cw, ch, "Router.decide(Event)")
y += ch + Inches(0.12)
chip(slide, cx, y, cw, ch, "Detector registry")
y += ch + Inches(0.15)
lane_w = Inches(4.6)
lane_gap = cw - 2 * lane_w
lx1, lx2 = cx, cx + lane_w + lane_gap
textbox(slide, lx1, y, lane_w, Inches(0.3), "LOG PATH", size=9, color=MUTED, font=MONO)
textbox(slide, lx2, y, lane_w, Inches(0.3), "FILE PATH", size=9, color=MUTED, font=MONO)
ly = y + Inches(0.32)
for text in ["log_ngram", "log_hst", "log_sgd"]:
    chip(slide, lx1, ly, lane_w, Inches(0.42), text)
    ly += Inches(0.48)
ly = y + Inches(0.32)
for text, kw in [("file_static", {}), ("file_ml_or_rf", {}),
                  ("borderline score? escalate", dict(dashed=True, text_color=MUTED)),
                  ("file_sandbox_signal, bwrap isolated", {}),
                  ("file_qsvc, quantum SVM", dict(border=QUANTUM, text_color=QUANTUM))]:
    chip(slide, lx2, ly, lane_w, Inches(0.42), text, **kw)
    ly += Inches(0.48)
y = max(ly, y + Inches(0.32) + 3 * Inches(0.48)) + Inches(0.1)
chip(slide, cx, y, cw, ch, "Confidence-weighted fusion")
y += ch + Inches(0.1)
chip(slide, cx, y, cw, ch, "Verdict: score, confidence, suspicious or malicious")
y += ch + Inches(0.1)
chip(slide, cx, y, cw, ch, "Encrypted quarantine, Fernet + HMAC (confidence 0.5+)",
     border=RUST, text_color=RUST)

# 4. Detection models and results (table) ------------------------------------
slide = body_slide("Detectors", "04 / 12", "Detection models and results")
det_rows = [
    ("Detector", "Type", "Trained on", "Measured"),
    ("log_ngram / log_hst", "Unsupervised, online", "none, calibrates live", "non-constant scores, verified"),
    ("log_sgd", "Supervised, online", "LogHub HDFS_v1", "acc 0.9953 · AUC 0.9823"),
    ("file_static", "Heuristic", "none", "entropy/header triage, always first"),
    ("file_ml_or_rf", "Supervised (LightGBM)", "EMBER2018", "acc 0.9491 · AUC 0.9891"),
    ("file_sandbox_signal", "Static enrichment", "n/a", "read-only, bwrap-isolated, never executes"),
    ("file_qsvc", "Hybrid classical + quantum", "EMBER2018 subset", "acc 0.6575 · AUC 0.6911"),
]
tbl = slide.shapes.add_table(len(det_rows), 4, MARGIN, Inches(1.1), CONTENT_W, Inches(3.6)).table
widths = [Inches(2.6), Inches(2.8), Inches(2.8), Inches(3.6)]
for i, w in enumerate(widths):
    tbl.columns[i].width = w
for r_i, row in enumerate(det_rows):
    for c_i, val in enumerate(row):
        cell = tbl.cell(r_i, c_i)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURFACE if r_i == 0 else PAPER
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(11 if r_i == 0 else 12)
                r.font.name = MONO if r_i == 0 else SANS
                r.font.bold = (r_i == 0)
                r.font.color.rgb = MUTED if r_i == 0 else (QUANTUM if row[0] == "file_qsvc" else INK)
textbox(slide, MARGIN, Inches(5.0), CONTENT_W, Inches(0.5),
        "The quantum stage's AUC is documented in-repo as “modest,” the honest weak "
        "link. Nothing here is invented for the pitch.", size=12, color=MUTED, font=SANS)

# 5. Decision fusion logic ---------------------------------------------------
slide = body_slide("Fusion", "05 / 12", "Decision fusion logic")
textbox(slide, MARGIN, Inches(1.1), CONTENT_W, Inches(1.0),
        "score = sum(score * confidence) / sum(confidence), verdict flips at score >= 0.5. "
        "The router short-circuits early: score >= 0.75 at confidence >= 0.5 is malicious, "
        "stop. Score <= 0.25 at confidence >= 0.5 is benign, stop.",
        size=13, color=INK, font=SANS)
cw2 = Inches(2.6)
cx2 = MARGIN
y2 = Inches(2.3)
chip(slide, cx2, y2, cw2, Inches(0.5), "file_static")
chip(slide, cx2 + cw2 + Inches(0.3), y2, cw2, Inches(0.5), "file_ml_or_rf")
chip(slide, cx2 + 2 * (cw2 + Inches(0.3)), y2, cw2 + Inches(0.6), Inches(0.5),
     "borderline? 0.45-0.65", dashed=True, text_color=MUTED)
y2 += Inches(0.8)
textbox(slide, cx2, y2, Inches(2.3), Inches(0.3), "NO", size=10, color=MUTED, font=MONO)
chip(slide, cx2, y2 + Inches(0.35), Inches(2.6), Inches(0.5), "Return fused verdict")
textbox(slide, cx2 + Inches(3.2), y2, Inches(2.3), Inches(0.3), "YES", size=10, color=MUTED, font=MONO)
chip(slide, cx2 + Inches(3.2), y2 + Inches(0.35), Inches(2.6), Inches(0.5), "Sandbox signal")
chip(slide, cx2 + Inches(6.1), y2 + Inches(0.35), Inches(2.8), Inches(0.5),
     "file_qsvc, quantum SVM", border=QUANTUM, text_color=QUANTUM)
chip(slide, cx2 + Inches(3.2), y2 + Inches(1.0), Inches(2.6), Inches(0.5), "Return fused verdict")
textbox(slide, MARGIN, Inches(4.7), CONTENT_W, Inches(0.6),
        "The quantum SVM is a tie-breaker, not a gate every file passes through. That's "
        "what makes it real-time compatible.", size=12, color=MUTED, font=SANS)

# 6. Sandbox and quarantine ---------------------------------------------------
slide = body_slide("Security design", "06 / 12", "Sandbox and quarantine")
half_w = (CONTENT_W - Inches(0.5)) / 2
textbox(slide, MARGIN, Inches(1.2), half_w, Inches(0.35), "SANDBOX", size=12, color=MUTED, font=SANS, bold=True)
textbox(slide, MARGIN, Inches(1.65), half_w, Inches(2.4),
        "bwrap with --unshare-all --die-with-parent --new-session --clearenv, read-only "
        "binds, private tmpfs. A real functional probe, not a presence check, falls back "
        "to rlimit-only isolation if bwrap can't actually sandbox. Never executes the "
        "sample by default.", size=13, color=INK, font=SANS)
textbox(slide, MARGIN + half_w + Inches(0.5), Inches(1.2), half_w, Inches(0.35), "QUARANTINE",
        size=12, color=MUTED, font=SANS, bold=True)
textbox(slide, MARGIN + half_w + Inches(0.5), Inches(1.65), half_w, Inches(2.4),
        "Fernet encryption, key at chmod 0600, HMAC-SHA256 integrity check on the on-disk "
        "database. Write-verify-then-delete-original ordering: quarantining can never "
        "silently lose data.", size=13, color=INK, font=SANS)

# 7. Training datasets ---------------------------------------------------------
slide = body_slide("Datasets", "07 / 12", "Training datasets")
textbox(slide, MARGIN, Inches(1.15), CONTENT_W, Inches(1.2),
        "Early planning scoped a wide field of options: CIC-IDS2017/2018, UNSW-NB15, "
        "MalwareBazaar, theZoo, alongside LogHub and EMBER. What's actually wired into "
        "training today is the narrower, concrete set below, both real and both already "
        "producing the measured numbers in this pitch.", size=13, color=INK, font=SANS)
textbox(slide, MARGIN, Inches(2.6), half_w, Inches(0.5), "LogHub HDFS_v1", size=18, color=RUST, font=MONO)
textbox(slide, MARGIN, Inches(3.15), half_w, Inches(0.5), "real distributed-system logs, feeds log_sgd",
        size=11, color=MUTED, font=SANS)
textbox(slide, MARGIN + half_w + Inches(0.5), Inches(2.6), half_w, Inches(0.5), "EMBER2018",
        size=18, color=RUST, font=MONO)
textbox(slide, MARGIN + half_w + Inches(0.5), Inches(3.15), half_w, Inches(0.5),
        "real malware PE features, feeds file_ml_or_rf + file_qsvc", size=11, color=MUTED, font=SANS)
textbox(slide, MARGIN, Inches(4.1), CONTENT_W, Inches(0.7),
        "The wider list stays the roadmap for broadening coverage past logs and PE files. "
        "It's not a claim about what's trained now.", size=12, color=MUTED, font=SANS, italic=True)

# 8. Desktop client features ---------------------------------------------------
slide = body_slide("Desktop client", "08 / 12", "Desktop client features")
items = [
    ("Dashboard", "Real cumulative counters: files scanned, lines checked, monitoring uptime. Not placeholders."),
    ("Scan", "On-demand file scan, wired to the live orchestrator."),
    ("Log analysis", "Background tailing and manual line-check, raw tail and flagged-lines table."),
    ("Quarantine", "Encrypted store with list, restore, and delete."),
    ("Adjustable UI", "Resizable columns and sidebar, both persisted across restarts."),
    ("Theme", "Live light/dark toggle, no restart required."),
]
y3 = Inches(1.15)
for tag, desc in items:
    textbox(slide, MARGIN, y3, Inches(1.7), Inches(0.5), tag, size=11, color=MUTED, font=MONO)
    textbox(slide, MARGIN + Inches(1.8), y3, CONTENT_W - Inches(1.8), Inches(0.5), desc, size=13, color=INK, font=SANS)
    hline(slide, MARGIN, y3 + Inches(0.55), CONTENT_W)
    y3 += Inches(0.72)

# 9. Testing and CI pipeline ----------------------------------------------------
slide = body_slide("Engineering rigor", "09 / 12", "Testing and CI pipeline")
stats = [("216", "tests passing"), ("~80%", "real coverage"), ("6", "CI gates")]
sx = MARGIN
for val, label in stats:
    textbox(slide, sx, Inches(1.2), Inches(2.4), Inches(0.55), val, size=26, color=RUST, font=MONO)
    textbox(slide, sx, Inches(1.8), Inches(2.4), Inches(0.4), label, size=11, color=MUTED, font=SANS)
    sx += Inches(2.7)
textbox(slide, MARGIN, Inches(2.6), CONTENT_W, Inches(1.0),
        "Every push runs a style checker, ruff, mypy, bandit, pip-audit, and pytest with a "
        "coverage floor, plus a separate gitleaks secret-scan and a Docker build with a "
        "health-check job.", size=13, color=INK, font=SANS)
textbox(slide, MARGIN, Inches(3.7), CONTENT_W, Inches(0.8),
        "For a tool whose entire job is telling you what to trust, its own supply chain "
        "and code quality have to hold up to the same scrutiny.", size=12, color=MUTED, font=SANS)

# 10. Challenges faced during development ---------------------------------------
slide = body_slide("Challenges", "10 / 12", "Challenges faced during development")
items = [
    ("Training safety", "Real malware features, without risking the machines doing the training. Training ran in isolated VMs, process killed the moment it strayed into dangerous territory."),
    ("Offline bar", "Treated as load-bearing, not aspirational: no scan, verdict, or quarantine action depends on reaching anything off-box."),
    ("Threat research", "Consumer malware reporting doesn't map cleanly onto a high-assurance threat model, so detectors were built around generalizable primitives, not one narrow attack corpus."),
    ("No cloud budget", "Every training run, every CI check, every test in the 216-test suite runs on ordinary developer hardware. That's itself evidence the system doesn't secretly need infrastructure a real deployment wouldn't have."),
]
y3 = Inches(1.15)
for tag, desc in items:
    textbox(slide, MARGIN, y3, Inches(1.9), Inches(0.5), tag, size=11, color=MUTED, font=MONO)
    textbox(slide, MARGIN + Inches(2.0), y3, CONTENT_W - Inches(2.0), Inches(1.0), desc, size=12.5, color=INK, font=SANS)
    y3 += Inches(1.28)

# 11. Limitations and future work -------------------------------------------------
slide = body_slide("Honest limits", "11 / 12", "Limitations and future work")
boxes = [
    ("RANSOMWARE MODEL", "No ransomware-labeled dataset or classifier yet. The log-anomaly pipeline is the right primitive, positioned for it, not claiming it's done."),
    ("SHELVED PIPELINE", "A separate TLS/encrypted-traffic-analysis pipeline exists in the repo but was never wired to real data or the live backend. Deliberately shelved, not silently abandoned."),
    ("CAPE SANDBOX ADAPTER", "Planned, documented as a NotImplementedError stub. External by design, since CAPE is GPLv3-licensed."),
]
y3 = Inches(1.15)
for k, v in boxes:
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, y3, CONTENT_W, Inches(1.15))
    box.fill.solid(); box.fill.fore_color.rgb = SURFACE
    box.line.color.rgb = HAIRLINE; box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(14); tf.margin_top = Pt(8)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run(); r0.text = k; r0.font.size = Pt(10); r0.font.name = MONO; r0.font.color.rgb = QUANTUM; r0.font.bold = True
    p1 = tf.add_paragraph()
    r1 = p1.add_run(); r1.text = v; r1.font.size = Pt(12.5); r1.font.name = SANS; r1.font.color.rgb = INK
    y3 += Inches(1.35)

# 12. Demo walkthrough -----------------------------------------------------------
slide = body_slide("Demo", "12 / 12", "Demo walkthrough")
textbox(slide, MARGIN, Inches(1.2), CONTENT_W, Inches(1.6),
        "The demo video covers, in order: a live file scan and verdict, background log "
        "tailing with real anomaly flags landing in the UI, an encrypted-quarantine "
        "restore and delete cycle, and the adjustable UI with a live theme toggle, all "
        "against the real running backend, no mocked data.", size=14, color=INK, font=SANS)
textbox(slide, MARGIN, Inches(2.9), CONTENT_W, Inches(0.5),
        "Full shot-by-shot script: hackathon/demo-video-script.md", size=12, color=MUTED, font=MONO)

# 13. Close -----------------------------------------------------------------------
slide = new_slide()
textbox(slide, MARGIN, Inches(3.0), CONTENT_W, Inches(1.1), "RAKSHAK", size=52, color=INK,
        font=SERIF, bold=True, align=PP_ALIGN.CENTER)
textbox(slide, MARGIN, Inches(3.9), CONTENT_W, Inches(0.6),
        "Hybrid classical and quantum threat detection, built real, scored honest.",
        size=16, color=MUTED, font=SERIF, italic=True, align=PP_ALIGN.CENTER)

out = Path(__file__).parent / "RAKSHAK-deck.pptx"
prs.save(out)
print(f"wrote {out} ({len(prs.slides)} slides)")
